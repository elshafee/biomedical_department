import os
import shutil
from services.word_editor import replace_placeholders
from docx import Document
from PIL import Image

# Dummy image
img = Image.new('RGB', (100, 100), color = 'red')
img.save('dummy.jpg')

shutil.copy("template_essic.docx", "test_out_combined.docx")

replacements = {
    "{{CODE_NUMBER}}": "0002 AIE",
    "{{SEND_TO}}": " ",
    "{{SUBJECT}}": " ",
    "{{STACK_HOLDER}}": " ",
    "{{POSITION}}": " ",
    "{{BODY_TEXT}}": "Hello\n{{TABLE_1}}\nWorld\n{{IMAGE_1}}\nEnd",
    "{{SENDER}}": " ",
    "{{SENDER_TOP}}": " ",
    "{{AIE}}": " "
}

table_data = [["Col 1", "Col 2"], ["Val 1", "Val 2"]]

replace_placeholders("template_essic.docx", "test_out_combined.docx", replacements, table_data=table_data, image_paths=['dummy.jpg'])

doc = Document("test_out_combined.docx")
print(f"Number of tables found: {len(doc.tables)}")
print(f"Number of inline shapes (images) found: {len(doc.inline_shapes)}")

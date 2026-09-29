import os
from services.word_editor import replace_placeholders
import shutil

shutil.copy("template_essic.docx", "test_output.docx")

sender_name_used = "أحمد خالد"
sender_position = "رئيس قسم"
sender_top = f"\u202A{sender_name_used} / {sender_position}\u202C"

replacements = {
    "{{CODE_NUMBER}}": "0001",
    "{{SEND_TO}}": " ",
    "{{SUBJECT}}": " ",
    "{{STACK_HOLDER}}": " ",
    "{{POSITION}}": " ",
    "{{BODY_TEXT}}": " ",
    "{{SENDER}}": " ",
    "{{SENDER_TOP}}": sender_top,
    "{{AIE}}": " ",
}

replace_placeholders("template_essic.docx", "test_output.docx", replacements)

from docx import Document
doc = Document("test_output.docx")
print("Output SENDER line:")
print([p.text for p in doc.paragraphs if "الراســــــــل" in p.text])

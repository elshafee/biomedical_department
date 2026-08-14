import re

with open("services/word_editor.py", "r", encoding="utf-8") as f:
    content = f.read()

# 1. Update signature
content = content.replace(
    'def replace_placeholders(input_path: str, output_path: str, replacements: dict, remove_manager_sig: bool = False):',
    'def replace_placeholders(input_path: str, output_path: str, replacements: dict, remove_manager_sig: bool = False, table_data: list = None, image_paths: list = None):'
)

# 2. Update the modified_body section
old_body_logic = """                    # Set the first part to the current paragraph
                    paragraph.runs[0].text = parts[0]
                    for run in paragraph.runs[1:]:
                        run.text = ""
                        
                    # Apply formatting to current paragraph
                    paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.JUSTIFY
                    fmt = paragraph.paragraph_format
                    fmt.space_before = Pt(0)
                    fmt.space_after = Pt(0)
                    fmt.line_spacing = 1.0

                    pPr = paragraph._element.get_or_add_pPr()
                    bidi = OxmlElement('w:bidi')
                    bidi.set(qn('w:val'), '1')
                    pPr.append(bidi)
                    
                    # Add subsequent paragraphs
                    current_p = paragraph
                    for part in parts[1:]:
                        if not part.strip():
                            continue
                        new_p_element = OxmlElement('w:p')
                        current_p._element.addnext(new_p_element)
                        new_p = Paragraph(new_p_element, current_p._parent)
                        
                        run = new_p.add_run(part)
                        if paragraph.runs and paragraph.runs[0]._element.rPr is not None:
                            run._element.append(copy.deepcopy(paragraph.runs[0]._element.rPr))
                        
                        new_p.alignment = WD_PARAGRAPH_ALIGNMENT.JUSTIFY
                        fmt = new_p.paragraph_format
                        fmt.space_before = Pt(0)
                        fmt.space_after = Pt(0)
                        fmt.line_spacing = 1.0

                        pPr = new_p._element.get_or_add_pPr()
                        bidi = OxmlElement('w:bidi')
                        bidi.set(qn('w:val'), '1')
                        pPr.append(bidi)
                        
                        current_p = new_p"""

new_body_logic = """                    # Clear original paragraph text completely
                    for run in paragraph.runs:
                        run.text = ""
                    
                    current_p = paragraph
                    first_text_added = False
                    
                    for part in parts:
                        if not part.strip():
                            continue
                            
                        if "{{TABLE_1}}" in part and table_data:
                            try:
                                from docx.enum.table import WD_TABLE_ALIGNMENT
                                new_table = doc.add_table(rows=len(table_data), cols=len(table_data[0]))
                                new_table.style = 'Table Grid'
                                new_table.alignment = WD_TABLE_ALIGNMENT.CENTER
                                tblPr = new_table._element.xpath('w:tblPr')
                                if tblPr:
                                    bidiVisual = OxmlElement('w:bidiVisual')
                                    tblPr[0].append(bidiVisual)
                                for r_idx, row_data in enumerate(table_data):
                                    for c_idx, cell_value in enumerate(row_data):
                                        cell = new_table.cell(r_idx, c_idx)
                                        cell.text = str(cell_value)
                                        for p in cell.paragraphs:
                                            p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
                                            pPr = p._element.get_or_add_pPr()
                                            bidi = OxmlElement('w:bidi')
                                            bidi.set(qn('w:val'), '1')
                                            pPr.append(bidi)
                                tbl_xml = new_table._element
                                tbl_xml.getparent().remove(tbl_xml)
                                current_p._element.addnext(tbl_xml)
                                
                                dummy_p_xml = OxmlElement('w:p')
                                tbl_xml.addnext(dummy_p_xml)
                                current_p = Paragraph(dummy_p_xml, current_p._parent)
                            except Exception as e:
                                print("Table error:", e)
                            continue
                            
                        if "{{IMAGE_1}}" in part and image_paths:
                            try:
                                from docx.shared import Inches
                                for img_path in image_paths:
                                    img_p_xml = OxmlElement('w:p')
                                    current_p._element.addnext(img_p_xml)
                                    img_p = Paragraph(img_p_xml, current_p._parent)
                                    img_p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
                                    run = img_p.add_run()
                                    run.add_picture(img_path, width=Inches(5.5))
                                    current_p = img_p
                            except Exception as e:
                                print("Image error:", e)
                            continue
                            
                        # Normal text part
                        if not first_text_added:
                            new_p = paragraph
                            first_text_added = True
                        else:
                            new_p_element = OxmlElement('w:p')
                            current_p._element.addnext(new_p_element)
                            new_p = Paragraph(new_p_element, current_p._parent)
                            
                        run = new_p.add_run(part)
                        if paragraph.runs and paragraph.runs[0]._element.rPr is not None:
                            run._element.append(copy.deepcopy(paragraph.runs[0]._element.rPr))
                        
                        new_p.alignment = WD_PARAGRAPH_ALIGNMENT.JUSTIFY
                        fmt = new_p.paragraph_format
                        fmt.space_before = Pt(0)
                        fmt.space_after = Pt(0)
                        fmt.line_spacing = 1.0

                        pPr = new_p._element.get_or_add_pPr()
                        bidi = OxmlElement('w:bidi')
                        bidi.set(qn('w:val'), '1')
                        pPr.append(bidi)
                        
                        current_p = new_p"""

content = content.replace(old_body_logic, new_body_logic)

with open("services/word_editor.py", "w", encoding="utf-8") as f:
    f.write(content)

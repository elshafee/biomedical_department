import os
import shutil
from services.word_editor import replace_placeholders
from docx import Document

# Create a dummy image
from PIL import Image
img = Image.new('RGB', (100, 100), color = 'red')
img.save('dummy.jpg')

shutil.copy("template_essic.docx", "test_out3.docx")

replacements = {
    "{{CODE_NUMBER}}": "0003 AIE",
    "{{SEND_TO}}": " ",
    "{{SUBJECT}}": " ",
    "{{STACK_HOLDER}}": " ",
    "{{POSITION}}": " ",
    "{{BODY_TEXT}}": "Hello\n{{IMAGE_1}}\nWorld",
    "{{SENDER}}": " ",
    "{{SENDER_TOP}}": " ",
    "{{AIE}}": " "
}

replace_placeholders("template_essic.docx", "test_out3.docx", replacements, image_paths=['dummy.jpg'])

doc = Document("test_out3.docx")
print("Done")
